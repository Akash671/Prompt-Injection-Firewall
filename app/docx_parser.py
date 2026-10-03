from io import BytesIO

from docx import Document


def extract_docx_text(data: bytes) -> str:

    document = Document(
        BytesIO(data)
    )

    parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    return "\n".join(parts)